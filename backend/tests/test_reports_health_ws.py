import asyncio
import io
import json
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook
from starlette.testclient import TestClient
from starlette.websockets import WebSocketDisconnect

from app.core.security import create_access_token
from app.main import app
from app.services import heartbeat, realtime

from .conftest import add_device, make_user
from .test_sla_rollup import seed


async def test_relatorio_resumo(db, alice, client):
    a = await add_device(db, alice.id, "A", "10.0.0.1")
    await add_device(db, alice.id, "B", "10.0.0.2")
    await seed(db, a.id, hours=24, end=datetime.now(UTC))
    body = (await client.get("/api/reports/summary?period=24h")).json()
    rows = {r["name"]: r for r in body["devices"]}
    assert (
        body["summary"]["devices"] == 2
        and rows["A"]["uptime_pct"] == 83.333
        and rows["B"]["uptime_pct"] is None
    )
    assert body["summary"]["avg_uptime_pct"] == 83.333  # ignora quem não tem medição
    assert (await client.get("/api/reports/summary?period=99d")).status_code == 422


async def test_relatorio_isolado_por_usuario(db, alice, bob_client):
    await add_device(db, alice.id, "Meu", "10.0.0.1")
    assert (await bob_client.get("/api/reports/summary")).json()["devices"] == []


async def test_csv_bloqueia_injecao_de_formula(db, alice, client):
    await add_device(db, alice.id, '=HYPERLINK("http://evil","x")', "10.0.0.1", location="@SUM(A1)")
    r = await client.get("/api/reports/export.csv")
    text = r.content.decode("utf-8-sig")
    assert (
        r.headers["content-type"].startswith("text/csv") and "attachment" in r.headers["content-disposition"]
    )
    assert "'=HYPERLINK" in text and "'@SUM" in text
    assert not any(line.startswith("=") for line in text.splitlines())


async def test_xlsx_valido_e_sem_formula(db, alice, client):
    await add_device(db, alice.id, "+cmd|' /C calc'!A0", "10.0.0.1")
    r = await client.get("/api/reports/export.xlsx")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value == "Nome" and ws["A2"].value.startswith("'+cmd") and ws["A2"].data_type != "f"


async def test_health_liveness(client):
    assert (await client.get("/health")).json() == {"status": "ok"}


async def test_readiness_sem_heartbeat_fica_degradado_sem_503(client):
    r = await client.get("/health/ready")
    assert r.status_code == 200 and r.json()["status"] == "degraded"
    assert (
        r.json()["checks"]["database"]["status"] == "ok"
        and r.json()["checks"]["monitor"]["status"] == "no_heartbeat"
    )


async def test_readiness_com_worker_vivo(client):
    await heartbeat.write({"devices": 3, "online": 2, "offline": 1})
    body = (await client.get("/health/ready")).json()
    assert body["status"] == "ok" and body["checks"]["monitor"]["offline"] == 1


async def test_readiness_detecta_worker_parado(client, redis):
    old = (datetime.now(UTC) - timedelta(minutes=10)).isoformat()
    await redis.set(heartbeat.KEY, json.dumps({"at": old, "devices": 1}))
    body = (await client.get("/health/ready")).json()
    assert body["checks"]["monitor"]["status"] == "stale" and body["status"] == "degraded"


async def test_readiness_503_quando_banco_cai(client, monkeypatch):
    class Broken:
        def connect(self):
            raise ConnectionError("db down")

    monkeypatch.setattr("app.api.routes.health.get_engine", lambda: Broken())
    r = await client.get("/health/ready")
    assert r.status_code == 503 and r.json()["status"] == "error"


def _ws_checks(token_alice: str, token_bob: str):
    tc = TestClient(app, base_url="http://test")
    results = {}
    with pytest.raises(WebSocketDisconnect) as no_auth, tc.websocket_connect("/ws"):
        pass
    results["no_auth"] = no_auth.value.code
    with (
        pytest.raises(WebSocketDisconnect) as bad_origin,
        tc.websocket_connect(
            "/ws", headers={"origin": "https://evil.example", "cookie": f"orbnoc_session={token_alice}"}
        ),
    ):
        pass
    results["bad_origin"] = bad_origin.value.code
    with tc.websocket_connect(
        "/ws", headers={"cookie": f"orbnoc_session={token_alice}", "origin": "http://testserver"}
    ) as ws:
        results["alice"] = ws.receive_json()
        ws.send_text("ping")
        results["pong"] = ws.receive_text()
    with tc.websocket_connect("/ws", headers={"authorization": f"Bearer {token_bob}"}) as ws:
        results["bob"] = ws.receive_json()
    return results


async def test_websocket_auth_origem_snapshot_e_isolamento(db):
    alice, bob = await make_user(db, "alice"), await make_user(db, "bob")
    await add_device(db, alice.id, "Do Alice", "10.0.0.1")
    res = await asyncio.to_thread(_ws_checks, create_access_token(alice.id), create_access_token(bob.id))
    assert res["no_auth"] == 1008 and res["bad_origin"] == 1008 and res["pong"] == "pong"
    assert [d["name"] for d in res["alice"]["data"]] == ["Do Alice"] and res["alice"]["type"] == "devices"
    assert res["bob"]["data"] == []


class DummyWS:
    def __init__(self):
        self.received = []

    async def send_json(self, msg):
        self.received.append(msg)


async def test_redis_repassa_apenas_ao_dono_do_dado(redis):
    mine, other = DummyWS(), DummyWS()
    realtime.manager.connect(1, mine)
    realtime.manager.connect(2, other)
    task = asyncio.create_task(realtime.manager.listen())
    try:
        for _ in range(60):
            await realtime.publish("event", 1, {"message": "olá"})
            await asyncio.sleep(0.05)
            if mine.received:
                break
    finally:
        task.cancel()
        realtime.manager.disconnect(1, mine)
        realtime.manager.disconnect(2, other)
    assert mine.received[0] == {"type": "event", "data": {"message": "olá"}} and other.received == []


async def test_socket_quebrado_e_removido(redis):
    class Broken:
        async def send_json(self, _):
            raise RuntimeError("closed")

    ws = Broken()
    realtime.manager.connect(9, ws)
    await realtime.manager.send_to_user(9, {"x": 1})
    assert realtime.manager.connection_count == 0
