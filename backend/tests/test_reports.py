import io
from datetime import UTC, datetime, timedelta

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.models import Device, Metric, User
from app.services.sla import rollup_hourly


async def seed(alice, sessionmaker, name="Core"):
    async with sessionmaker() as s:
        uid = (await s.scalar(select(User).where(User.username == "alice"))).id
        d = Device(
            user_id=uid,
            name=name,
            ip="10.0.0.1",
            location="DC1",
            status="online",
            sla_threshold_ms=80,
        )
        s.add(d)
        await s.flush()
        now = datetime.now(UTC) - timedelta(minutes=5)
        for i in range(10):
            s.add(
                Metric(
                    device_id=d.id,
                    recorded_at=now - timedelta(seconds=10 * i),
                    ok=i != 0,
                    latency=20.0 if i else None,
                )
            )
        await s.commit()
        await rollup_hourly(s, datetime.now(UTC), hours=3)
        await s.commit()


async def test_summary_and_sla(alice, sessionmaker):
    await seed(alice, sessionmaker)
    s = (await alice.get("/api/reports/summary?window=24h")).json()
    row = s["devices"][0]
    assert row["uptime_pct"] == 90.0 and row["avg_latency"] == 20.0 and row["sample_count"] == 10
    assert s["average_uptime_pct"] == 90.0 and s["total_devices"] == 1
    assert (await alice.get("/api/sla?window=30d")).json()["devices"][0]["uptime_pct"] == 90.0
    assert (await alice.get("/api/reports/summary?window=1y")).status_code == 422


async def test_uptime_series(alice, sessionmaker):
    await seed(alice, sessionmaker)
    series = (await alice.get("/api/uptime-series?hours=6")).json()
    assert (
        series
        and series[-1]["uptime"] == 90.0
        and series[-1]["online"] == 1
        and series[-1]["offline"] == 0
    )


async def test_csv_export_and_formula_injection(alice, sessionmaker):
    await seed(alice, sessionmaker, name='=HYPERLINK("http://evil","x")')
    r = await alice.get("/api/reports/export?format=csv&window=24h")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert "attachment" in r.headers["content-disposition"]
    lines = r.text.lstrip("\ufeff").splitlines()
    assert lines[0].startswith("Dispositivo,Host")
    assert lines[1].startswith("\"'=HYPERLINK") or lines[1].startswith("'=HYPERLINK")


async def test_xlsx_export(alice, sessionmaker):
    await seed(alice, sessionmaker)
    r = await alice.get("/api/reports/export?format=xlsx&window=7d")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]][:2] == ["Dispositivo", "Host"]
    assert ws["A2"].value == "Core" and ws["E2"].value == 90.0


async def test_reports_are_per_user(alice, bob_client, sessionmaker):
    await seed(alice, sessionmaker)
    assert (await bob_client.get("/api/reports/summary")).json()["devices"] == []
    assert (await bob_client.get("/api/uptime-series")).json() == []


async def test_reports_require_auth(client):
    for url in ("/api/reports/summary", "/api/reports/export", "/api/uptime-series", "/api/sla"):
        assert (await client.get(url)).status_code == 401
